import os
import pandas as pd
from pypdf import PdfReader
from docx import Document
from openpyxl import load_workbook
from config import FILES_ROOT


def _safe_path(relative_path: str) -> str:
    full_path = os.path.abspath(os.path.join(FILES_ROOT, relative_path))
    if not full_path.startswith(os.path.abspath(FILES_ROOT)):
        raise ValueError("Access outside the Alfred files folder is not allowed.")
    return full_path


def _truncate(text: str, limit: int = 6000) -> str:
    if len(text) > limit:
        return text[:limit] + "\n...[content truncated]"
    return text


def read_pdf(relative_path: str) -> str:
    try:
        target = _safe_path(relative_path)
    except ValueError as e:
        return str(e)

    if not os.path.isfile(target):
        return f"File '{relative_path}' does not exist."

    try:
        reader = PdfReader(target)
        text = "\n".join(page.extract_text() or "" for page in reader.pages)
    except Exception as e:
        return f"Couldn't read PDF: {e}"

    if not text.strip():
        return "This PDF appears to have no extractable text (it may be scanned/image-based)."
    return _truncate(text)


def read_word_doc(relative_path: str) -> str:
    try:
        target = _safe_path(relative_path)
    except ValueError as e:
        return str(e)

    if not os.path.isfile(target):
        return f"File '{relative_path}' does not exist."

    try:
        doc = Document(target)
        text = "\n".join(p.text for p in doc.paragraphs)
    except Exception as e:
        return f"Couldn't read Word document: {e}"

    return _truncate(text) if text.strip() else "This document appears to be empty."


def read_excel(relative_path: str, sheet_name: str = None) -> str:
    try:
        target = _safe_path(relative_path)
    except ValueError as e:
        return str(e)

    if not os.path.isfile(target):
        return f"File '{relative_path}' does not exist."

    try:
        wb = load_workbook(target, data_only=True)
        sheet = wb[sheet_name] if sheet_name else wb.active

        rows = []
        for row in sheet.iter_rows(values_only=True):
            rows.append(" | ".join(str(c) if c is not None else "" for c in row))
        text = "\n".join(rows)
    except Exception as e:
        return f"Couldn't read Excel file: {e}"

    available_sheets = ", ".join(wb.sheetnames)
    header = f"Sheet: {sheet.title} (available sheets: {available_sheets})\n\n"
    return _truncate(header + text)


def read_csv(relative_path: str) -> str:
    try:
        target = _safe_path(relative_path)
    except ValueError as e:
        return str(e)

    if not os.path.isfile(target):
        return f"File '{relative_path}' does not exist."

    try:
        df = pd.read_csv(target)
        summary = f"Columns: {', '.join(df.columns)}\nRows: {len(df)}\n\n"
        preview = df.head(30).to_string(index=False)
    except Exception as e:
        return f"Couldn't read CSV: {e}"

    return _truncate(summary + preview)